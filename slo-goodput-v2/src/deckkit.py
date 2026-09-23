# -*- coding: utf-8 -*-
"""轻量级 PPT 绘制工具层：统一的配色、字体、卡片、箭头、表格。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from copy import deepcopy

# ---------------------------------------------------------------- 设计变量
SW, SH = 13.333, 7.5          # 16:9
ML = 0.78                      # 左边距
CW = SW - 2 * ML               # 内容宽度

INK   = "14213D"   # 主标题 深海军蓝
BODY  = "3F4A63"   # 正文
MUTED = "8B95A8"   # 次要信息
FAINT = "AEB6C4"
LINE  = "E2E8F0"   # 分割线
PANEL = "F5F7FA"   # 浅面板
PANEL2= "EEF2F8"
WHITE = "FFFFFF"

BLUE  = "2563EB"; BLUE_L  = "E8EFFD"; BLUE_D = "1D4ED8"
TEAL  = "0E9488"; TEAL_L  = "E0F5F2"
AMBER = "C2810C"; AMBER_L = "FDF4E3"
RED   = "D0454C"; RED_L   = "FCECEC"
VIO   = "7A5AF8"; VIO_L   = "EFEBFE"

EA = "Microsoft YaHei"   # 中文字体
LA = "Segoe UI"          # 西文字体
MONO = "Consolas"


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------- 字体
def _apply_font(run, size, bold, color, latin, ea, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor.from_string(color)
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", ea), ("a:cs", ea)):
        el = rPr.find(qn(tag))
        if el is None:
            el = parse_xml('<a:%s xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
                           % tag.split(":")[1])
            rPr.append(el)
        el.set("typeface", face)


def textbox(slide, x, y, w, h, anchor="t"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, text, size=13, bold=False, color=BODY, align="l", space_before=0,
         space_after=0, line=1.28, first=False, latin=LA, ea=EA, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    _apply_font(r, size, bold, color, latin, ea, italic)
    return p


def rich(tf, chunks, size=13, color=BODY, align="l", space_before=0, space_after=0,
         line=1.28, first=False):
    """chunks: [(text, {bold:.., color:.., size:..}), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for t, o in chunks:
        r = p.add_run()
        r.text = t
        _apply_font(r, o.get("size", size), o.get("bold", False), o.get("color", color),
                    o.get("latin", LA), o.get("ea", EA), o.get("italic", False))
    return p


def txt(slide, x, y, w, h, text, size=13, bold=False, color=BODY, align="l",
        anchor="t", line=1.28, latin=LA, ea=EA):
    tb, tf = textbox(slide, x, y, w, h, anchor)
    para(tf, text, size, bold, color, align, line=line, first=True, latin=latin, ea=ea)
    return tb


# ---------------------------------------------------------------- 形状
def _nofill(sh):
    sh.fill.background()


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0, radius=None,
          shadow=False):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        _nofill(sh)
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(lw)
    if not shadow:
        sh.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.text_frame.word_wrap = True
    sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(0.12)
    sh.text_frame.margin_top = sh.text_frame.margin_bottom = Inches(0.06)
    return sh


def rect(slide, x, y, w, h, fill=PANEL, line=None, lw=1.0):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, lw)


def card(slide, x, y, w, h, fill=WHITE, line=LINE, lw=1.0, radius=0.07):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, radius)


def pill(slide, x, y, w, h, text, size=11.5, bold=True, fill=BLUE_L, fc=BLUE,
         line=None, lw=1.0):
    sh = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, 0.5)
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size, bold, fc, "c", first=True, line=1.0)
    return sh


def box_text(sh, lines, anchor="m", align="c"):
    """lines: [(text, size, bold, color)]"""
    tf = sh.text_frame
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, (t, s, b, c) in enumerate(lines):
        para(tf, t, s, b, c, align, first=(i == 0), line=1.2,
             space_before=0 if i == 0 else 2)


def arrow(slide, x1, y1, x2, y2, color=FAINT, lw=1.5, dash=False, head="triangle",
          size="med"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = parse_xml('<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>')
        ln.append(d)
    if head:
        w = {"sm": "sm", "med": "med", "lg": "lg"}[size]
        t = parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                      'type="%s" w="%s" len="%s"/>' % (head, w, w))
        ln.append(t)
    c.shadow.inherit = False
    return c


def chevron(slide, x, y, w, h, text, fill=BLUE_L, fc=INK, size=12, bold=True):
    sh = shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill, None)
    box_text(sh, [(text, size, bold, fc)])
    return sh


def tri_down(slide, cx, y, w=0.2, h=0.14, color=FAINT):
    sh = shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - w / 2, y, w, h, color, None)
    sh.rotation = 180
    return sh


def hline(slide, x, y, w, color=LINE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                   Inches(x + w), Inches(y))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    c.shadow.inherit = False
    return c


def vline(slide, x, y, h, color=LINE, lw=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                   Inches(x), Inches(y + h))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    c.shadow.inherit = False
    return c


# ---------------------------------------------------------------- 页面骨架
FOOTER = "大模型推理服务：SLO 与 Goodput  ·  技术分享"


def header(slide, kicker, title, sub=None, accent=BLUE):
    if kicker:
        rect(slide, ML, 0.52, 0.045, 0.2, accent)
        txt(slide, ML + 0.14, 0.5, 8.0, 0.26, kicker, 11, True, accent, line=1.0)
    txt(slide, ML, 0.82, CW, 0.52, title, 27, True, INK, line=1.05)
    y = 1.42
    if sub:
        txt(slide, ML, y, CW - 0.3, 0.3, sub, 13, False, MUTED, line=1.2)
        y += 0.42
    hline(slide, ML, y + 0.02, CW)
    return y + 0.3


def page_no(slide, n):
    txt(slide, ML, SH - 0.52, 7.0, 0.24, FOOTER, 8.5, False, FAINT, line=1.0)
    txt(slide, SW - ML - 1.2, SH - 0.52, 1.2, 0.24, "%02d" % n, 9.5, True, FAINT, "r", line=1.0)


def takeaway(slide, y, text, accent=BLUE, fill=BLUE_L, h=0.62, size=14.5, icon="核心结论"):
    """页面底部的结论条"""
    card(slide, ML, y, CW, h, fill, None, radius=0.12)
    rect(slide, ML, y + 0.08, 0.05, h - 0.16, accent)
    tb, tf = textbox(slide, ML + 0.26, y, CW - 0.5, h, "m")
    rich(tf, [(icon + "  ｜  ", {"bold": True, "color": accent, "size": size - 2.5}),
              (text, {"bold": True, "color": INK, "size": size})], first=True, line=1.2)


def section(prs, num, title, desc, n, accent=BLUE, faint=PANEL2):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, 0.16, SH, accent)
    txt(s, ML + 0.1, 1.45, 5.0, 2.2, num, 120, True, faint, line=0.9)
    rect(s, ML + 0.22, 3.62, 0.5, 0.045, accent)
    txt(s, ML + 0.2, 3.9, 9.5, 0.6, title, 34, True, INK, line=1.05)
    txt(s, ML + 0.2, 4.72, 8.6, 0.8, desc, 14.5, False, MUTED, line=1.45)
    page_no(s, n)
    return s


def grid_table(slide, x, y, w, col_fr, rows, row_h=0.42, head_h=0.44,
               accent=INK, head_fill=PANEL2, zebra=PANEL, size=11.5,
               head_size=11.5, align=None, first_bold=True, first_color=INK):
    """手绘表格：rows[0] 为表头。col_fr 为列宽比例。"""
    total = sum(col_fr)
    xs, acc = [], x
    for fr in col_fr:
        xs.append((acc, w * fr / total))
        acc += w * fr / total
    align = align or ["l"] + ["c"] * (len(col_fr) - 1)
    # 表头
    rect(slide, x, y, w, head_h, head_fill)
    for j, ((cx, cwid), cell) in enumerate(zip(xs, rows[0])):
        txt(slide, cx + 0.14, y + (head_h - 0.22) / 2, cwid - 0.28, 0.24, cell,
            head_size, True, accent, align[j], line=1.0)
    yy = y + head_h
    for i, row in enumerate(rows[1:]):
        if i % 2 == 1:
            rect(slide, x, yy, w, row_h, zebra)
        for j, ((cx, cwid), cell) in enumerate(zip(xs, row)):
            b = first_bold and j == 0
            c = first_color if j == 0 else BODY
            txt(slide, cx + 0.14, yy + (row_h - 0.22) / 2, cwid - 0.28, 0.24, cell,
                size, b, c, align[j], line=1.0)
        yy += row_h
        hline(slide, x, yy, w, LINE, 0.75)
    hline(slide, x, y, w, LINE, 0.75)
    return yy


def formula(slide, x, y, w, h, text, fill=PANEL, fc=INK, size=15, line_color=None,
            bold=True):
    sh = card(slide, x, y, w, h, fill, line_color, 1.0, 0.1)
    box_text(sh, [(text, size, bold, fc)])
    return sh
