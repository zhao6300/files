# -*- coding: utf-8 -*-
"""近似渲染 pptx -> PNG，用于自查版式（非精确排版，仅用于发现溢出/重叠）。"""
import sys, os
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

DPI = 100
FONT_R = "/usr/share/fonts/google-noto-cjk/NotoSansCJKsc-Regular.otf"
FONT_B = "/usr/share/fonts/google-noto-cjk/NotoSansCJKsc-Bold.otf"
_cache = {}


def font(size_pt, bold):
    key = (round(size_pt * 2) / 2, bold)
    if key not in _cache:
        px = max(6, int(round(size_pt * DPI / 72.0)))
        _cache[key] = ImageFont.truetype(FONT_B if bold else FONT_R, px)
    return _cache[key]


def emu_px(v):
    return int(round(Emu(v).inches * DPI))


def solid_rgb(fmt):
    try:
        if fmt.type is not None and fmt.type == 1:  # MSO_FILL.SOLID
            c = fmt.fore_color
            if c.type == 1:  # RGB
                return "#" + str(c.rgb)
    except Exception:
        pass
    return None


def line_rgb(sh):
    try:
        ln = sh.line
        if ln.fill.type == 1:
            return "#" + str(ln.color.rgb), max(1, int(round((ln.width.pt if ln.width else 1) * DPI / 72.0)))
    except Exception:
        pass
    return None, 0


def wrap(text, f, maxw, draw):
    """按像素宽度折行，支持中英混排。"""
    out, cur = [], ""
    for ch in text:
        if ch == "\n":
            out.append(cur); cur = ""; continue
        t = cur + ch
        if draw.textlength(t, font=f) > maxw and cur:
            out.append(cur); cur = ch
        else:
            cur = t
    out.append(cur)
    return out


BOXES = []


def draw_text(draw, sh, x, y, w, h, warns, name):
    tf = sh.text_frame if sh.has_text_frame else None
    if tf is None:
        return
    ml = emu_px(tf.margin_left or 0); mr = emu_px(tf.margin_right or 0)
    mt = emu_px(tf.margin_top or 0); mb = emu_px(tf.margin_bottom or 0)
    bx, by = x + ml, y + mt
    bw, bh = max(4, w - ml - mr), max(4, h - mt - mb)
    lines = []   # (text, font, color, align, height)
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            lines.append((None, None, None, None, 6))
            continue
        size = runs[0].font.size.pt if runs[0].font.size else 18
        bold = bool(runs[0].font.bold)
        col = "#000000"
        try:
            if runs[0].font.color and runs[0].font.color.type == 1:
                col = "#" + str(runs[0].font.color.rgb)
        except Exception:
            pass
        f = font(size, bold)
        ls = p.line_spacing if p.line_spacing else 1.2
        if not isinstance(ls, float) and not isinstance(ls, int):
            ls = 1.2
        lh = size * 1.2 * float(ls) * DPI / 72.0
        al = str(p.alignment) if p.alignment is not None else "LEFT (1)"
        sb = (p.space_before.pt if p.space_before else 0) * DPI / 72.0
        sa = (p.space_after.pt if p.space_after else 0) * DPI / 72.0
        txt = "".join(r.text for r in runs)
        if sb:
            lines.append((None, None, None, None, sb))
        for seg in wrap(txt, f, bw, draw):
            lines.append((seg, f, col, al, lh))
        if sa:
            lines.append((None, None, None, None, sa))
    total = sum(l[4] for l in lines)
    anchor = str(sh.text_frame.vertical_anchor)
    if "MIDDLE" in anchor:
        cy = by + (bh - total) / 2
    elif "BOTTOM" in anchor:
        cy = by + bh - total
    else:
        cy = by
    if total > bh + 3:
        warns.append("溢出(高): %s 需要%.0fpx 实际%.0fpx | %s" % (name, total, bh,
                     (lines[0][0] or "")[:24]))
    for seg, f, col, al, lh in lines:
        if seg is None:
            cy += lh; continue
        tw = draw.textlength(seg, font=f)
        if "CENTER" in al:
            tx = bx + (bw - tw) / 2
        elif "RIGHT" in al:
            tx = bx + bw - tw
        else:
            tx = bx
        asc = f.getmetrics()[0]
        draw.text((tx, cy + (lh - asc) / 2 - 1), seg, font=f, fill=col)
        BOXES.append((tx, cy + 2, tx + tw, cy + lh - 2, name, seg))
        cy += lh


def render(path, outdir, only=None):
    prs = Presentation(path)
    W = emu_px(prs.slide_width); H = emu_px(prs.slide_height)
    os.makedirs(outdir, exist_ok=True)
    allw = []
    for idx, slide in enumerate(prs.slides, 1):
        if only and idx not in only:
            continue
        img = Image.new("RGB", (W, H), "#FFFFFF")
        d = ImageDraw.Draw(img)
        warns = []
        del BOXES[:]
        for sh in slide.shapes:
            try:
                x, y = emu_px(sh.left), emu_px(sh.top)
                w, h = emu_px(sh.width), emu_px(sh.height)
            except Exception:
                continue
            st = sh.shape_type
            if sh.element.tag.endswith('}cxnSp'):
                xf = sh.element.find('.//' + qn('a:xfrm'))
                fh = xf is not None and xf.get('flipH') == '1'
                fv = xf is not None and xf.get('flipV') == '1'
                x1, x2 = (x + w, x) if fh else (x, x + w)
                y1, y2 = (y + h, y) if fv else (y, y + h)
                c, lw = line_rgb(sh)
                d.line([x1, y1, x2, y2], fill=c or "#999999", width=max(1, lw))
                continue
            fill = solid_rgb(sh.fill) if hasattr(sh, "fill") else None
            lc, lw = line_rgb(sh) if st != MSO_SHAPE_TYPE.TEXT_BOX else (None, 0)
            shp = str(getattr(sh, "shape_type", ""))
            auto = None
            try:
                auto = str(sh.auto_shape_type)
            except Exception:
                pass
            if fill or lc:
                if auto and "OVAL" in auto:
                    d.ellipse([x, y, x + w, y + h], fill=fill, outline=lc, width=max(1, lw))
                elif auto and "ROUNDED" in auto:
                    d.rounded_rectangle([x, y, x + w, y + h], radius=min(12, h // 4),
                                        fill=fill, outline=lc, width=max(1, lw))
                elif auto and "TRIANGLE" in auto:
                    d.polygon([(x + w / 2, y + h), (x, y), (x + w, y)], fill=fill, outline=lc)
                else:
                    d.rectangle([x, y, x + w, y + h], fill=fill, outline=lc, width=max(1, lw))
            draw_text(d, sh, x, y, w, h, warns, "p%d/#%s" % (idx, sh.shape_id))
            if x < -2 or y < -2 or x + w > W + 2 or y + h > H + 2:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    warns.append("越界: p%d %s (%.2f,%.2f,%.2f,%.2f)" %
                                 (idx, sh.text_frame.text[:18], x / DPI, y / DPI,
                                  (x + w) / DPI, (y + h) / DPI))
        # 文字重叠检测（同一形状内的行不比较）
        for i in range(len(BOXES)):
            for j in range(i + 1, len(BOXES)):
                a, b = BOXES[i], BOXES[j]
                if a[4] == b[4]:
                    continue
                ox = min(a[2], b[2]) - max(a[0], b[0])
                oy = min(a[3], b[3]) - max(a[1], b[1])
                if ox > 6 and oy > 5:
                    warns.append("文字重叠: p%d 『%s』 × 『%s』 (%dx%dpx)" %
                                 (idx, a[5][:16], b[5][:16], ox, oy))
        del BOXES[:]
        img.save(os.path.join(outdir, "s%02d.png" % idx))
        for wn in warns:
            allw.append(wn)
    print("\n".join(allw) if allw else "无警告")


if __name__ == "__main__":
    only = set(int(a) for a in sys.argv[1:]) if len(sys.argv) > 1 else None
    here = os.path.dirname(os.path.abspath(__file__))
    render(os.path.join(here, "..", "LLM-Inference-SLO-Goodput-v2.pptx"),
           os.path.join(here, "..", "preview"), only)
